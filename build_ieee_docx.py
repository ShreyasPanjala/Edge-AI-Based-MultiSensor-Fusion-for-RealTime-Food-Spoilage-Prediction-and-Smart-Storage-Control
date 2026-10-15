"""Convert conference_paper_ieee.md to an authentic IEEE two-column Word document."""

import re
from pathlib import Path
import docx
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

MD_PATH = Path("conference_paper_ieee.md")
DOCX_PATH = Path("conference_paper_ieee.docx")


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Set custom thin borders on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{top or "none"}" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="{left or "none"}" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="{bottom or "none"}" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="{right or "none"}" w:sz="4" w:space="0" w:color="000000"/>\n'
        f"</w:tcBorders>"
    )
    tcPr.append(tcBorders)


def set_cell_shading(cell, color_hex="F2F2F2"):
    """Set light background shading on a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def create_ieee_document():
    doc = docx.Document()

    # --- SECTION 1: Single column for Title and Author block ---
    s1 = doc.sections[0]
    s1.top_margin = Inches(0.75)
    s1.bottom_margin = Inches(1.0)
    s1.left_margin = Inches(0.625)
    s1.right_margin = Inches(0.625)

    # Base style setup
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Times New Roman"
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)

    # Paper Title (24 pt, Bold, Centered)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.line_spacing = 1.15
    run_title = p_title.add_run(
        "A Zero-Allocation TinyML Stacking Hybrid Architecture for Real-Time Volatile Fruit Spoilage Detection on Edge Microcontrollers"
    )
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(24)
    run_title.bold = True

    # Author Table (3 authors centered across 3 columns)
    auth_table = doc.add_table(rows=1, cols=3)
    auth_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    authors_data = [
        ("Author 1", "dept. name of organization", "name of organization", "City, Country", "author1@institution.edu"),
        ("Author 2", "dept. name of organization", "name of organization", "City, Country", "author2@institution.edu"),
        ("Author 3", "dept. name of organization", "name of organization", "City, Country", "author3@institution.edu"),
    ]
    for col_idx, (name, dept, org, loc, email) in enumerate(authors_data):
        cell = auth_table.rows[0].cells[col_idx]
        cell.width = Inches(2.4)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0

        r_name = p.add_run(f"{name}\n")
        r_name.font.name = "Times New Roman"
        r_name.font.size = Pt(11)
        r_name.bold = True

        r_dept = p.add_run(f"{dept}\n{org}\n{loc}\n")
        r_dept.font.name = "Times New Roman"
        r_dept.font.size = Pt(10)
        r_dept.italic = True

        r_mail = p.add_run(email)
        r_mail.font.name = "Times New Roman"
        r_mail.font.size = Pt(9)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)

    # --- SECTION 2: Continuous Break, Two Columns for Abstract, Index Terms & Body ---
    s2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
    s2.top_margin = Inches(0.75)
    s2.bottom_margin = Inches(1.0)
    s2.left_margin = Inches(0.625)
    s2.right_margin = Inches(0.625)

    # Configure exact 2-column layout in XML
    sectPr = s2._sectPr
    cols = sectPr.xpath("./w:cols")
    if cols:
        cols[0].set(qn("w:num"), "2")
        cols[0].set(qn("w:space"), "360")  # 360 dxa = 0.25 inches spacing
        cols[0].set(qn("w:equalWidth"), "1")

    # Read markdown source content
    raw_md = MD_PATH.read_text(encoding="utf-8")

    # Abstract
    abs_match = re.search(r"### Abstract\s*\n(.*?)\n\n\*\*Index Terms\*\*—(.*?)\n", raw_md, re.DOTALL)
    if abs_match:
        abstract_text = abs_match.group(1).strip()
        index_terms_text = abs_match.group(2).strip()

        p_abs = doc.add_paragraph()
        p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_abs.paragraph_format.space_before = Pt(0)
        p_abs.paragraph_format.space_after = Pt(4)
        p_abs.paragraph_format.line_spacing = 1.0
        r_lbl = p_abs.add_run("Abstract—")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(9)
        r_lbl.bold = True
        r_lbl.italic = True

        r_txt = p_abs.add_run(abstract_text)
        r_txt.font.name = "Times New Roman"
        r_txt.font.size = Pt(9)
        r_txt.bold = True

        p_idx = doc.add_paragraph()
        p_idx.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_idx.paragraph_format.space_before = Pt(0)
        p_idx.paragraph_format.space_after = Pt(12)
        p_idx.paragraph_format.line_spacing = 1.0
        r_ilbl = p_idx.add_run("Index Terms—")
        r_ilbl.font.name = "Times New Roman"
        r_ilbl.font.size = Pt(9)
        r_ilbl.bold = True
        r_ilbl.italic = True

        r_itxt = p_idx.add_run(index_terms_text)
        r_itxt.font.name = "Times New Roman"
        r_itxt.font.size = Pt(9)

    # Process Body Sections
    # Split content after Index Terms
    after_index = raw_md.split("**Index Terms**", 1)[1]
    # Remove header separator if present
    lines = [ln.rstrip() for ln in after_index.split("\n")][2:]

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    table_caption = ""

    def flush_code():
        nonlocal code_lines
        if not code_lines:
            return
        p_code = doc.add_paragraph()
        p_code.paragraph_format.space_before = Pt(4)
        p_code.paragraph_format.space_after = Pt(6)
        p_code.paragraph_format.line_spacing = 1.0
        run_c = p_code.add_run("\n".join(code_lines))
        run_c.font.name = "Consolas"
        run_c.font.size = Pt(7.5)
        run_c.font.color.rgb = RGBColor(40, 40, 40)
        code_lines = []

    def flush_table():
        nonlocal table_lines, table_caption
        if not table_lines:
            return
        # Parse markdown table
        headers = [c.strip() for c in table_lines[0].split("|") if c.strip()]
        rows = []
        for r_str in table_lines[2:]:  # skip separator
            cells = [c.strip() for c in r_str.split("|")[1:-1]]
            if cells:
                rows.append(cells)

        if headers and rows:
            # Table Title in Small Caps
            if table_caption:
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(8)
                p_cap.paragraph_format.space_after = Pt(2)
                r_tc = p_cap.add_run(table_caption.upper())
                r_tc.font.name = "Times New Roman"
                r_tc.font.size = Pt(8)
                r_tc.font.small_caps = True
                r_tc.bold = True

            t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER

            # Header row
            for c_idx, h_text in enumerate(headers):
                cell = t.rows[0].cells[c_idx]
                set_cell_borders(cell, top="single", bottom="single")
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(h_text.replace("**", ""))
                r.font.name = "Times New Roman"
                r.font.size = Pt(7.5)
                r.bold = True

            # Data rows
            for r_idx, r_data in enumerate(rows):
                is_last = r_idx == len(rows) - 1
                row_cells = t.rows[r_idx + 1].cells
                for c_idx, val in enumerate(r_data[: len(headers)]):
                    cell = row_cells[c_idx]
                    set_cell_borders(cell, top="none", bottom="single" if is_last else "none")
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    r = p.add_run(val.replace("**", ""))
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(7.5)
                    if "**" in val:
                        r.bold = True

            # Spacer after table
            p_tsp = doc.add_paragraph()
            p_tsp.paragraph_format.space_before = Pt(0)
            p_tsp.paragraph_format.space_after = Pt(6)

        table_lines = []
        table_caption = ""

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # Handle Code blocks
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code()
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(lines[i])
            i += 1
            continue

        # Handle Table Captions
        if line.startswith("TABLE "):
            table_caption = line
            i += 1
            continue

        # Handle Markdown Tables
        if line.startswith("|") and line.endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            in_table = False
            flush_table()

        # Handle Section Headings (## Heading)
        if line.startswith("## "):
            sec_title = line.replace("## ", "").strip()
            p_sec = doc.add_paragraph()
            p_sec.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sec.paragraph_format.space_before = Pt(12)
            p_sec.paragraph_format.space_after = Pt(4)
            p_sec.paragraph_format.keep_with_next = True
            r_sec = p_sec.add_run(sec_title.upper())
            r_sec.font.name = "Times New Roman"
            r_sec.font.size = Pt(10)
            r_sec.font.small_caps = True
            r_sec.bold = True
            i += 1
            continue

        # Handle Subsection Headings (### Heading)
        if line.startswith("### "):
            subsec_title = line.replace("### ", "").strip()
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_sub.paragraph_format.space_before = Pt(8)
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.keep_with_next = True
            r_sub = p_sub.add_run(subsec_title)
            r_sub.font.name = "Times New Roman"
            r_sub.font.size = Pt(10)
            r_sub.italic = True
            i += 1
            continue

        # Handle Display Equations
        if line.startswith("$$") and line.endswith("$$") and len(line) > 4:
            eq_text = line[2:-2].strip()
            p_eq = doc.add_paragraph()
            p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_eq.paragraph_format.space_before = Pt(4)
            p_eq.paragraph_format.space_after = Pt(4)
            r_eq = p_eq.add_run(eq_text)
            r_eq.font.name = "Cambria Math"
            r_eq.font.size = Pt(9.5)
            r_eq.italic = True
            i += 1
            continue

        # Handle List Items
        if re.match(r"^(\d+\.|\*|-)\s+", line):
            clean_item = re.sub(r"^(\d+\.|\*|-)\s+", "", line)
            p_li = doc.add_paragraph()
            p_li.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_li.paragraph_format.left_indent = Inches(0.18)
            p_li.paragraph_format.space_before = Pt(0)
            p_li.paragraph_format.space_after = Pt(2)
            p_li.paragraph_format.line_spacing = 1.0

            # Bullet run
            bullet_sym = "• " if not line[0].isdigit() else line.split()[0] + " "
            r_b = p_li.add_run(bullet_sym)
            r_b.font.name = "Times New Roman"
            r_b.font.size = Pt(9.5)
            r_b.bold = True

            # Parse bold/regular runs
            parts = re.split(r"(\*\*.*?\*\*)", clean_item)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p_li.add_run(part[2:-2])
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9.5)
                    r.bold = True
                else:
                    r = p_li.add_run(part)
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9.5)
            i += 1
            continue

        # Handle References
        if line.startswith("[") and "]" in line and len(line) > 5 and line[1].isdigit():
            p_ref = doc.add_paragraph()
            p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_ref.paragraph_format.left_indent = Inches(0.25)
            p_ref.paragraph_format.first_line_indent = Inches(-0.25)
            p_ref.paragraph_format.space_before = Pt(0)
            p_ref.paragraph_format.space_after = Pt(3)
            p_ref.paragraph_format.line_spacing = 1.0
            r_ref = p_ref.add_run(line)
            r_ref.font.name = "Times New Roman"
            r_ref.font.size = Pt(8)
            i += 1
            continue

        # Handle Regular Paragraphs
        if line and not line.startswith("---"):
            p_body = doc.add_paragraph()
            p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_body.paragraph_format.space_before = Pt(0)
            p_body.paragraph_format.space_after = Pt(4)
            p_body.paragraph_format.line_spacing = 1.0
            p_body.paragraph_format.first_line_indent = Inches(0.14)

            parts = re.split(r"(\*\*.*?\*\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p_body.add_run(part[2:-2])
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10)
                    r.bold = True
                else:
                    r = p_body.add_run(part)
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10)

        i += 1

    # Final flush if still in table or code
    flush_table()
    flush_code()

    doc.save(DOCX_PATH)
    print(f"Successfully generated IEEE Word document: {DOCX_PATH.resolve()}")


if __name__ == "__main__":
    create_ieee_document()
